"""
Regression: a .docx letterhead, and a text box, rode out in the approved file.

:func:`batch.write_approved_docx` re-scans the redacted document and refuses to
write if anything identifying survives -- its docstring calls this "the same bar
as the text path". That bar was not being met for two places text can hide.

The re-scan read the body and its tables only. A letterhead is not in the body:
the clinic name, clinician, address, postcode and phone sit in a **header or
footer**. Nor is a floating **text box** -- python-docx's paragraph walk does not
descend into a ``w:txbxContent``. Text in either place was invisible to
detection, so it never entered the approved map and was never redacted, and it
was invisible to the sweep that is meant to be the last line of defence. The
approved .docx was written with all of it intact.

Both are now swept, so such a document is refused rather than written. Refusing
is the safe failure: the clinician keeps the original and is told why, and
nothing identifying reaches disk. Every name, number and address below is
fabricated.
"""

from __future__ import annotations

import io

import docx
import pytest
from docx.oxml import parse_xml

from carescribe.core import batch, deidentify, ingest

BODY_PATIENT = "Wilhelmina Featherstonehaugh"
HEADER_CLINICIAN = "Dr Archibald Pemberton-Fitzroy"
FOOTER_MRN = "ZZ1122334"
FOOTER_ADDRESS = "14 Rosemary Gardens, Bolton BL1 4AB"
FOOTER_PHONE = "01204 555123"
BOX_CLINICIAN = "Dr Millicent Thorncastle"
BOX_MRN = "ZZ7766554"
BOX_PHONE = "01204 555999"

# A VML text box. DrawingML shapes wrap their text in the same w:txbxContent
# tag, so covering this shape covers both.
_TEXT_BOX_XML = """
<w:pict xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"
        xmlns:v="urn:schemas-microsoft-com:vml">
  <v:shape><v:textbox><w:txbxContent>
    <w:p><w:r><w:t>{text}</w:t></w:r></w:p>
  </w:txbxContent></v:textbox></v:shape>
</w:pict>
"""


def _document(*, header: str = "", footer: str = "", text_box: str = "") -> bytes:
    """A .docx whose body is ordinary and whose extras are the thing under test."""
    document = docx.Document()
    document.add_paragraph(f"Patient: {BODY_PATIENT}")
    document.add_paragraph("Impression: stable, discharged on sertraline 50mg.")
    if text_box:
        run = document.add_paragraph().add_run()
        run._r.append(parse_xml(_TEXT_BOX_XML.format(text=text_box)))
    section = document.sections[0]
    if header:
        section.header.paragraphs[0].text = header
    if footer:
        section.footer.paragraphs[0].text = footer
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _approved_map(source_bytes: bytes, tmp_path) -> dict[str, str]:
    """Run the real flow the app runs: read the file, de-identify, build the map."""
    path = tmp_path / "source.docx"
    path.write_bytes(source_bytes)
    result = deidentify.deidentify(ingest.extract_text(str(path)))
    return {value: key for key, value in result.phi_map.items()}


def _write(source_bytes: bytes, tmp_path):
    return batch.write_approved_docx(
        "letterhead.docx",
        source_bytes,
        _approved_map(source_bytes, tmp_path),
        output_dir=tmp_path / "out",
    )


LETTERHEAD = {
    "header": f"St Elsewhere Clinic - {HEADER_CLINICIAN}",
    "footer": f"Ref {FOOTER_MRN} - {FOOTER_ADDRESS} - tel {FOOTER_PHONE}",
}
BOXED = {"text_box": f"Referrer {BOX_CLINICIAN}, MRN {BOX_MRN}, tel {BOX_PHONE}"}


# ==========================================================================
# The regression: the write is refused, and nothing lands
# ==========================================================================


@pytest.mark.parametrize(
    "extras, expected",
    [
        (LETTERHEAD, (FOOTER_PHONE, "BL1 4AB")),
        (BOXED, (BOX_PHONE, BOX_MRN)),
    ],
    ids=["letterhead", "text-box"],
)
def test_hidden_identifiers_refuse_the_write(extras, expected, tmp_path):
    """Before the fix this call succeeded and wrote every one of them to disk."""
    with pytest.raises(batch.BatchError, match="Refusing to write the Word file") as caught:
        _write(_document(**extras), tmp_path)
    for identifier in expected:
        assert identifier in str(caught.value)


@pytest.mark.parametrize("extras", [LETTERHEAD, BOXED], ids=["letterhead", "text-box"])
def test_nothing_reaches_disk_when_the_write_is_refused(extras, tmp_path):
    """The point of a refusal is that the leak never lands."""
    with pytest.raises(batch.BatchError):
        _write(_document(**extras), tmp_path)
    assert list((tmp_path / "out").glob("*.docx")) == []


# ==========================================================================
# Why the sweep has to be the thing that catches them
# ==========================================================================


def test_the_body_identifier_is_detected_normally(tmp_path):
    """The document is otherwise processed fine -- the extras are the gap."""
    assert BODY_PATIENT in _approved_map(_document(**LETTERHEAD), tmp_path)


@pytest.mark.parametrize(
    "extras, hidden",
    [
        (LETTERHEAD, (HEADER_CLINICIAN, FOOTER_MRN, FOOTER_PHONE)),
        (BOXED, (BOX_CLINICIAN, BOX_MRN, BOX_PHONE)),
    ],
    ids=["letterhead", "text-box"],
)
def test_hidden_identifiers_are_never_detected(extras, hidden, tmp_path):
    """Detection cannot see them, which is why the residual sweep must."""
    approved = _approved_map(_document(**extras), tmp_path)
    for identifier in hidden:
        assert identifier not in approved


# ==========================================================================
# Guards: a fix that refused every .docx would pass everything above
# ==========================================================================


def test_a_document_with_no_extras_still_writes(tmp_path):
    written = _write(_document(), tmp_path)
    assert written.exists()


def test_a_benign_header_and_footer_do_not_trip_the_scan(tmp_path):
    """A header with no identifier in it must not cost the clinician the file."""
    written = _write(
        _document(header="Clinical Record - Page 1", footer="Confidential"), tmp_path
    )
    assert written.exists()


def test_a_benign_text_box_does_not_trip_the_scan(tmp_path):
    written = _write(_document(text_box="Please file in the patient record."), tmp_path)
    assert written.exists()
