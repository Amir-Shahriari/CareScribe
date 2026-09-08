"""
Regression: the patient's name rode out in the .docx document properties.

Word fills a document's properties in from the machine that authored it, so a
clinical .docx arrives with ``dc:creator`` and ``cp:lastModifiedBy`` holding
clinician names and ``dc:title`` holding a heading like "Discharge summary for
<patient>". None of that is in the document *text*, so the reviewer never saw it
in the chip table and redaction never touched it -- and the approved file went
out with the patient's name in File > Properties after every visible trace of it
had been removed.

``write_approved_docx`` now blanks those fields, and the residual sweep reads the
property parts back afterwards, so a field the blanking list misses still refuses
the write instead of riding out. ``dcterms:created`` and ``dcterms:modified``
cannot be blanked -- the schema types them as timestamps -- so they are
normalised to a fixed instant, because when a note was authored is the same class
of fact as a service date.

Every name here is fabricated.
"""

from __future__ import annotations

import io
import zipfile

import docx
import pytest

from carescribe.core import batch

PATIENT = "Wilhelmina Featherstonehaugh"
AUTHOR = "Dr Archibald Pemberton-Fitzroy"
LAST_MODIFIED_BY = "Nurse Millicent Thorncastle"
TITLE = f"Discharge summary for {PATIENT}"
COMPANY = "St Elsewhere Clinic"

CORE = "docProps/core.xml"


def _document(**properties) -> bytes:
    """A .docx whose body is already clean and whose properties are not."""
    document = docx.Document()
    document.add_paragraph("Patient: [PATIENT]")
    document.add_paragraph("Impression: stable, discharged on sertraline 50mg.")
    for name, value in properties.items():
        setattr(document.core_properties, name, value)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _write(source_bytes: bytes, tmp_path):
    # The body is already de-identified, so the map only has to be non-empty --
    # write_approved_docx refuses an empty one, and this test is about metadata.
    return batch.write_approved_docx(
        "notes.docx", source_bytes, {"[PATIENT]": "[PATIENT]"}, output_dir=tmp_path / "out"
    )


def _xml_of(path) -> str:
    with zipfile.ZipFile(path) as archive:
        return " ".join(
            archive.read(name).decode("utf-8", "replace")
            for name in archive.namelist()
            if name.endswith(".xml")
        )


# ==========================================================================
# The regression
# ==========================================================================


@pytest.mark.parametrize(
    "field, value",
    [
        ("author", AUTHOR),
        ("last_modified_by", LAST_MODIFIED_BY),
        ("title", TITLE),
        ("subject", PATIENT),
        ("keywords", PATIENT),
        ("comments", f"Reviewed at {COMPANY}"),
        ("category", COMPANY),
    ],
)
def test_an_identifying_property_does_not_survive_into_the_approved_file(
    field, value, tmp_path
):
    """Each of these used to ride out verbatim in the written .docx."""
    written = _write(_document(**{field: value}), tmp_path)
    assert value not in _xml_of(written)


def test_the_patient_name_in_the_title_is_gone(tmp_path):
    """The worst case: the title names the patient the body no longer does."""
    written = _write(_document(title=TITLE, author=AUTHOR), tmp_path)
    xml = _xml_of(written)
    assert PATIENT not in xml
    assert AUTHOR not in xml


def test_authoring_timestamps_are_normalised_not_left_telling_the_date(tmp_path):
    """A note's authoring time is the same class of fact as a service date."""
    import datetime

    seen_on = datetime.datetime(2024, 3, 5, 14, 30)
    written = _write(_document(created=seen_on, modified=seen_on), tmp_path)
    with zipfile.ZipFile(written) as archive:
        core = archive.read(CORE).decode("utf-8")
    assert "2024-03-05" not in core
    assert batch._NEUTRAL_TIMESTAMP in core


# ==========================================================================
# The sweep backstops the blanking list
# ==========================================================================


def test_a_property_the_blanking_misses_still_refuses_the_write(tmp_path, monkeypatch):
    """The field list is a list, so it must not be the only thing standing here.

    With blanking disabled, an identifying property has to be caught by the
    residual sweep reading the property parts back.
    """
    monkeypatch.setattr(batch, "_strip_document_properties", lambda data: data)
    with pytest.raises(batch.BatchError, match="Refusing to write the Word file"):
        _write(_document(author=AUTHOR), tmp_path)
    assert list((tmp_path / "out").glob("*.docx")) == []


# ==========================================================================
# Guards: the file must still be a working .docx
# ==========================================================================


def test_the_written_file_is_still_a_valid_readable_docx(tmp_path):
    written = _write(_document(author=AUTHOR, title=TITLE), tmp_path)

    assert zipfile.ZipFile(written).testzip() is None
    reopened = docx.Document(str(written))
    assert [p.text for p in reopened.paragraphs if p.text.strip()][0] == "Patient: [PATIENT]"


def test_every_part_of_the_original_package_survives(tmp_path):
    """Only the two property parts are rewritten; everything else is copied."""
    source = _document(author=AUTHOR)
    written = _write(source, tmp_path)

    original = set(zipfile.ZipFile(io.BytesIO(source)).namelist())
    result = set(zipfile.ZipFile(written).namelist())
    assert original - result == set()


def test_a_document_with_no_properties_set_still_writes(tmp_path):
    assert _write(_document(), tmp_path).exists()
